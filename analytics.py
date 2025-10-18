import psycopg2
import pandas as pd
import matplotlib.pyplot as plt
import plotly.express as px
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule

# Подключение к PostgreSQL
conn = psycopg2.connect(
    dbname="students_db",
    user="postgres",
    password="postgres123",   
    host="localhost",
    port="5432"
)

# Загружаем данные из таблицы
query = "SELECT * FROM students_raw;"
df = pd.read_sql(query, conn)

# ---------- ГРАФИКИ ----------

#  Pie chart – распределение по полу
gender_counts = df['gender'].value_counts()
gender_counts.plot(kind='pie', autopct='%1.1f%%')
plt.title('Распределение студентов по полу')
plt.savefig('charts/pie_gender.png')
plt.close()

#  Bar chart – средние оценки по предметам
df[['math_score', 'reading_score', 'writing_score']].mean().plot(kind='bar')
plt.title('Средние оценки по предметам')
plt.savefig('charts/bar_subjects.png')
plt.close()

#  Bar chart по этническим группам
group_avg = df.groupby('race_ethnicity')[['math_score', 'reading_score', 'writing_score']].mean()
group_avg.plot(kind='bar')
plt.title('Средние оценки по этническим группам')
plt.savefig('charts/bar_ethnicity.png')
plt.close()

#  Scatter – зависимость между математикой и письмом
plt.scatter(df['math_score'], df['writing_score'])
plt.xlabel('Math score')
plt.ylabel('Writing score')
plt.title('Зависимость между математикой и письмом')
plt.savefig('charts/scatter_math_write.png')
plt.close()

#  Histogram – распределение оценок по математике
df['math_score'].plot(kind='hist', bins=20)
plt.title('Распределение оценок по математике')
plt.savefig('charts/hist_math.png')
plt.close()

#  Line chart – добавим искусственные даты
df['exam_date'] = pd.date_range(start='2024-01-01', periods=len(df), freq='D')
avg_by_date = df.groupby('exam_date')['math_score'].mean()
avg_by_date.plot(kind='line')
plt.title('Изменение среднего балла по датам')
plt.savefig('charts/line_math_date.png')
plt.close()

# ---------- ИНТЕРАКТИВНЫЙ ГРАФИК ----------
fig = px.scatter(df, x='math_score', y='reading_score',
                 animation_frame='exam_date',
                 color='gender',
                 title='Изменение успеваемости по датам')
fig.write_html('charts/interactive_scatter.html')

# ---------- ЭКСПОРТ В EXCEL ----------
avg_scores = df.groupby('gender')[['math_score', 'reading_score', 'writing_score']].mean().reset_index()

with pd.ExcelWriter('exports/students_report.xlsx', engine='openpyxl') as writer:
    df.to_excel(writer, sheet_name='All Data', index=False)
    avg_scores.to_excel(writer, sheet_name='Average Scores', index=False)

wb = load_workbook('exports/students_report.xlsx')
for sheet in wb.sheetnames:
    ws = wb[sheet]
    rule = ColorScaleRule(start_type='min', start_color='FFAA0000',
                          mid_type='percentile', mid_value=50, mid_color='FFFFFF00',
                          end_type='max', end_color='FF00AA00')
    ws.conditional_formatting.add('C2:C100', rule)
wb.save('exports/students_report.xlsx')

print(" Аналитика завершена! Все графики и отчёты сохранены.")
